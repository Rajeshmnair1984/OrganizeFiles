#!/usr/bin/env python3
import os
import shutil
import argparse
import logging
import sys
import subprocess
import json
import hashlib
from pathlib import Path
from datetime import datetime
from typing import Optional, Dict, List, Any

# Setup logging first so it's available for install()
logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')

def install(package):
    """Install a package using pip."""
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", package])
        return True
    except subprocess.CalledProcessError as e:
        logging.error(f"Failed to install {package}: {e}")
        return False

# Auto-install dependencies
HAS_PDF = False
try:
    from pypdf import PdfReader
    HAS_PDF = True
except ImportError:
    print("Installing pypdf for PDF support...")
    if install("pypdf"):
        try:
            from pypdf import PdfReader
            HAS_PDF = True
        except ImportError: pass

HAS_DOCX = False
try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    print("Installing python-docx for DOCX support...")
    if install("python-docx"):
        try:
            from docx import Document
            HAS_DOCX = True
        except ImportError: pass

HAS_EXCEL = False
try:
    import openpyxl
    HAS_EXCEL = True
except ImportError:
    print("Installing openpyxl for Excel support...")
    if install("openpyxl"):
        try:
            import openpyxl
            HAS_EXCEL = True
        except ImportError: pass

HAS_PPTX = False
try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    print("Installing python-pptx for PowerPoint support...")
    if install("python-pptx"):
        try:
            from pptx import Presentation
            HAS_PPTX = True
        except ImportError: pass

# Default Configuration
DEFAULT_CATEGORIES = {
    "Documents": [".pdf", ".docx", ".doc", ".txt", ".xlsx", ".pptx", ".csv"],
    "Images": [".jpg", ".jpeg", ".png", ".gif", ".svg", ".webp", ".heic"],
    "Videos": [".mp4", ".mov", ".avi", ".mkv"],
    "Archives": [".zip", ".tar", ".gz", ".rar", ".7z"],
    "Music": [".mp3", ".wav", ".flac", ".m4a"],
    "Applications": [".dmg", ".pkg", ".app"],
    "Code": [".py", ".js", ".html", ".css", ".json", ".sh"],
}

DEFAULT_CONTENT_CATEGORIES = {
    "Financial": ["invoice", "bill", "receipt", "tax", "bank", "statement", "payment", "chequing", "banking", "account", "credit", "debit"],
    "Work": ["project", "report", "meeting", "resume", "proposal", "contract", "leads", "sales", "franchise"],
    "HR-Jobs": ["hr", "job", "hiring", "salary", "benefits", "employee"],
    "Legal": ["agreement", "legal", "court", "affidavit", "notary", "law", "signed"],
}

def load_config(config_path: Path) -> tuple[Dict, Dict]:
    """Loads configuration from a JSON file."""
    if config_path.exists():
        try:
            with config_path.open() as f:
                data = json.load(f)
                return (data.get("CATEGORIES", DEFAULT_CATEGORIES), 
                        data.get("CONTENT_CATEGORIES", DEFAULT_CONTENT_CATEGORIES))
        except Exception as e:
            logging.error(f"Failed to load config: {e}. Using defaults.")
    return DEFAULT_CATEGORIES, DEFAULT_CONTENT_CATEGORIES

def extract_text(file_path: Path) -> str:
    """Extracts text from various file formats."""
    ext = file_path.suffix.lower()
    text = ""
    try:
        if ext == ".txt":
            with file_path.open(errors='ignore') as f:
                text = f.read(100_000)
        elif ext == ".pdf" and HAS_PDF:
            reader = PdfReader(file_path)
            # Check metadata
            meta = reader.metadata
            if meta:
                text += f" {meta.get('/Title', '')} {meta.get('/Subject', '')} "
            # Extract from first few pages
            for page in reader.pages[:5]:
                text += page.extract_text() or ""
        elif ext == ".docx" and HAS_DOCX:
            doc = Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs[:50]])
        elif ext == ".xlsx" and HAS_EXCEL:
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
            for sheet in wb.worksheets[:3]:  # First 3 sheets
                for row in sheet.iter_rows(max_row=50, values_only=True):
                    text += " ".join([str(v) for v in row if v]) + " "
        elif ext == ".pptx" and HAS_PPTX:
            prs = Presentation(file_path)
            for slide in prs.slides[:10]:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + " "
    except Exception as e:
        logging.debug(f"Could not read content from {file_path.name}: {e}")
    return text.lower()

def classify(file_path: Path, content_categories: Dict) -> Optional[str]:
    """Classifies documents based on filename and content."""
    text = extract_text(file_path)
    content = f"{file_path.name} {text}".lower()
    
    scores = {cat: 0 for cat in content_categories}
    for cat, keywords in content_categories.items():
        for kw in keywords:
            if kw.lower() in content:
                # Weight filename hits higher than content hits
                scores[cat] += (content.count(kw.lower()))
                if kw.lower() in file_path.name.lower():
                    scores[cat] += 5 
                    
    if not any(scores.values()):
        return None
    return max(scores, key=lambda k: scores[k])

def file_hash(path: Path) -> str:
    """Returns MD5 hash of a file for duplicate content detection."""
    h = hashlib.md5()
    try:
        with path.open("rb") as f:
            while chunk := f.read(65536):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return ""

def build_hash_index(folder: Path) -> Dict[str, Path]:
    """Scan ALL files in folder recursively. Returns hash -> path of the oldest copy (the original)."""
    index: Dict[str, Path] = {}
    for f in folder.rglob("*"):
        if not f.is_file() or f.name.startswith('.'):
            continue
        h = file_hash(f)
        if not h:
            continue
        if h not in index:
            index[h] = f
        else:
            # Keep the older file as the "original"
            try:
                if f.stat().st_mtime < index[h].stat().st_mtime:
                    index[h] = f
            except OSError:
                pass
    return index

def get_safe_path(dest_dir: Path, filename: str) -> Path:
    """Handles duplicate filenames."""
    dest_path = dest_dir / filename
    if not dest_path.exists():
        return dest_path

    stem = Path(filename).stem
    suffix = Path(filename).suffix
    counter = 1
    while counter < 1000:
        new_name = f"{stem} ({counter}){suffix}"
        if not (dest_dir / new_name).exists():
            return dest_dir / new_name
        counter += 1
    return dest_dir / f"{stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{suffix}"

UNDO_LOG = Path.home() / ".organizer_undo.json"

def load_undo_log() -> List[Dict]:
    if UNDO_LOG.exists():
        try:
            with UNDO_LOG.open() as f:
                return json.load(f)
        except Exception:
            return []
    return []

def save_undo_log(entries: List[Dict]):
    with UNDO_LOG.open("w") as f:
        json.dump(entries, f, indent=2)

def undo_last_run():
    entries = load_undo_log()
    if not entries:
        print("No undo log found.")
        return

    last_session = entries[-1]["session"]
    session_entries = [e for e in entries if e["session"] == last_session]

    print(f"Undoing {len(session_entries)} moves from session {last_session}...")
    success, errors = 0, 0
    for entry in reversed(session_entries):
        src = Path(entry["src"])
        dst = Path(entry["dst"])
        try:
            if dst.exists():
                src.parent.mkdir(parents=True, exist_ok=True)
                shutil.move(str(dst), str(src))
                print(f"Restored: {dst.name} -> {src}")
                success += 1
            else:
                print(f"Skipped (not found): {dst}")
        except Exception as e:
            logging.error(f"Error restoring {dst}: {e}")
            errors += 1

    remaining = [e for e in entries if e["session"] != last_session]
    save_undo_log(remaining)
    print(f"\nRestored: {success}, Errors: {errors}")

def organize(folder: Path, categories: Dict, content_categories: Dict,
             dry_run: bool = True, by_date: bool = False, recursive: bool = False,
             session_id: str = "") -> tuple[int, int, Dict[str, int], List[Dict]]:
    moved, errors = 0, 0
    category_counts: Dict[str, int] = {}
    move_log: List[Dict] = []

    if not folder.exists():
        logging.warning(f"Folder not found: {folder}")
        return 0, 0, {}, []

    logging.info(f"--- Organizing: {folder} {'(DRY RUN)' if dry_run else ''} ---")

    # Scan ALL files in folder (including existing subfolders) to build duplicate index
    logging.info("Scanning all locations for duplicates...")
    hash_index = build_hash_index(folder)

    # Known category dirs — skip files already organized into them
    known_cats = set(categories.keys()) | {"Misc", "Duplicates"}

    items = list(folder.rglob("*") if recursive else folder.iterdir())

    for item in items:
        if item.is_dir() or item.name.startswith('.') or any(part.startswith('.') for part in item.parts):
            continue

        # Skip files already inside an organized subfolder
        try:
            rel = item.relative_to(folder)
            if rel.parts[0] in known_cats:
                continue
        except ValueError:
            pass

        # Check if this file is a duplicate of something already elsewhere in the folder tree
        h = file_hash(item)
        original = hash_index.get(h) if h else None
        is_dup = original is not None and original.resolve() != item.resolve()

        if is_dup:
            target_cat = "Duplicates"
            dest_dir = folder / "Duplicates"
        else:
            ext = item.suffix.lower()
            target_cat = "Misc"
            for cat, extensions in categories.items():
                if ext in extensions:
                    target_cat = cat
                    break

            dest_dir = folder / target_cat
            if target_cat == "Documents":
                sub_cat = classify(item, content_categories)
                dest_dir = dest_dir / (sub_cat if sub_cat else "Misc")

            if by_date:
                mtime = datetime.fromtimestamp(item.stat().st_mtime)
                dest_dir = dest_dir / str(mtime.year) / mtime.strftime("%m_%B")

        if not dry_run:
            dest_dir.mkdir(parents=True, exist_ok=True)

        final_dest = get_safe_path(dest_dir, item.name)

        try:
            rel_dest = dest_dir.relative_to(folder)
        except ValueError:
            rel_dest = dest_dir

        try:
            if dry_run:
                msg = f"[DRY RUN] {item.name} -> {rel_dest}/{final_dest.name}"
                if is_dup and original:
                    try:
                        orig_rel = original.relative_to(folder)
                    except ValueError:
                        orig_rel = original
                    msg += f"  [duplicate of {orig_rel}]"
                logging.info(msg)
            else:
                shutil.move(str(item), str(final_dest))
                log_msg = f"Moved: {item.name} -> {rel_dest}/{final_dest.name}"
                if is_dup and original:
                    try:
                        orig_rel = original.relative_to(folder)
                    except ValueError:
                        orig_rel = original
                    log_msg += f"  [duplicate of {orig_rel}]"
                logging.info(log_msg)
                move_log.append({"session": session_id, "src": str(item), "dst": str(final_dest)})
            moved += 1
            category_counts[target_cat] = category_counts.get(target_cat, 0) + 1
        except Exception as e:
            logging.error(f"Error moving {item.name}: {e}")
            errors += 1

    return moved, errors, category_counts, move_log

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Organize files logically.")
    parser.add_argument("folders", nargs="*", help="Folders to organize (defaults to Downloads and Documents)")
    parser.add_argument("--run", action="store_true", help="Execute the moves")
    parser.add_argument("--undo", action="store_true", help="Undo the last run")
    parser.add_argument("--by-date", action="store_true", help="Organize by year/month")
    parser.add_argument("--recursive", action="store_true", help="Organize subfolders too")
    parser.add_argument("--config", type=str, default="config.json", help="Path to config JSON")
    args = parser.parse_args()

    if args.undo:
        undo_last_run()
        sys.exit(0)

    folders = [Path(f) for f in args.folders] if args.folders else [Path.home() / "Downloads", Path.home() / "Documents"]
    categories, content_categories = load_config(Path(args.config))

    if not args.run:
        print("\nDRY RUN MODE: Use --run to execute.\n")

    session_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    t_moved, t_errors = 0, 0
    t_categories: Dict[str, int] = {}
    all_moves: List[Dict] = []

    for folder in folders:
        m, e, cats, moves = organize(
            folder, categories, content_categories,
            dry_run=not args.run, by_date=args.by_date,
            recursive=args.recursive, session_id=session_id
        )
        t_moved += m
        t_errors += e
        all_moves.extend(moves)
        for cat, count in cats.items():
            t_categories[cat] = t_categories.get(cat, 0) + count

    if args.run and all_moves:
        existing_log = load_undo_log()
        existing_log.extend(all_moves)
        save_undo_log(existing_log[-500:])  # cap at 500 entries
        print(f"Undo log saved to {UNDO_LOG}")

    print(f"\n--- Summary ---")
    print(f"Items Processed: {t_moved + t_errors}")
    print(f"Successful:      {t_moved}")
    print(f"Errors:          {t_errors}")
    if t_categories:
        print("\nBy Category:")
        for cat, count in sorted(t_categories.items(), key=lambda x: -x[1]):
            print(f"  {cat}: {count}")
